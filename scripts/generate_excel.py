from pathlib import Path
import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
rows = []
for p in sorted((ROOT / 'data/datasets').glob('*.yaml')):
    with open(p, encoding='utf-8') as f:
        r = yaml.safe_load(f) or {}
    e = r.get('embodiment') or {}
    s = r.get('scale') or {}
    a = r.get('access') or {}
    l = r.get('license') or {}
    links = r.get('links') or {}
    st = r.get('status') or {}
    legacy = r.get('legacy') or {}

    paper = links.get('paper') or legacy.get('论文/技术报告链接') or ''
    dataset = links.get('dataset') or a.get('download') or legacy.get('数据开源链接') or ''

    rows.append({
        'ID': r.get('id', ''),
        '数据集': r.get('name') or r.get('name_zh') or '',
        '发布年份': r.get('year', ''),
        '提出机构': '; '.join(r.get('organization', []) or []),
        '描述': r.get('description', ''),
        '数据来源': '; '.join(r.get('source', []) or []),
        '模态': '; '.join(r.get('modality', []) or []),
        '机器人本体': '; '.join(e.get('type', []) or []),
        '平台': '; '.join(e.get('platforms', []) or []),
        '任务': '; '.join(r.get('tasks', []) or []),
        '环境': '; '.join(r.get('environment', []) or []),
        '用途': '; '.join(r.get('application', []) or []),
        '规模备注': s.get('notes') or '',
        '数据格式': '; '.join(r.get('data_format', []) or []),
        '开放状态': a.get('availability') or '',
        '论文/技术报告链接': paper,
        '数据开源链接': dataset,
        '项目': links.get('project') or '',
        '代码': links.get('github') or '',
        'License': l.get('type') or '',
        'Last Verified': st.get('last_verified') or '',
    })

out = ROOT / 'exports/embodied-ai-datasets.xlsx'
out.parent.mkdir(exist_ok=True)
wb = Workbook()
ws = wb.active
ws.title = 'Dataset Catalog'
headers = list(rows[0].keys()) if rows else []
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])

header_fill = PatternFill('solid', fgColor='1F4E78')
header_font = Font(color='FFFFFF', bold=True)
for c in range(1, len(headers) + 1):
    cell = ws.cell(1, c)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

ws.freeze_panes = 'A2'
ws.auto_filter.ref = ws.dimensions

for colname in ['论文/技术报告链接', '数据开源链接', '项目', '代码']:
    if colname not in headers:
        continue
    c = headers.index(colname) + 1
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, c)
        if isinstance(cell.value, str) and cell.value.startswith(('http://', 'https://')):
            cell.hyperlink = cell.value
            cell.font = Font(color='0563C1', underline='single')

widths = [28, 30, 12, 32, 52, 25, 24, 24, 24, 28, 24, 28, 42, 24, 14, 48, 52, 45, 45, 16, 16]
for c, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.row_dimensions[1].height = 32
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 42
    for cell in ws[r]:
        cell.alignment = Alignment(vertical='top', wrap_text=True)

info = wb.create_sheet('说明')
info_rows = [
    ['字段', '说明'],
    ['论文/技术报告链接', '对应数据集论文、技术报告或预印本链接；如有 URL 则设置为可点击超链接。'],
    ['数据开源链接', '对应数据集官方下载页、Hugging Face/GitHub 等数据页面；如有 URL 则设置为可点击超链接。'],
    ['数据源', 'YAML Dataset Card；Excel 为自动生成导出，不是主数据源。'],
    ['当前数据规模', f'{len(rows)} 个 Dataset Cards；迁移源为 117 条原始记录，其中 Ego2Robot 有 1 条重复记录。'],
]
for row in info_rows:
    info.append(row)
for c in range(1, 3):
    info.cell(1, c).fill = header_fill
    info.cell(1, c).font = header_font
info.column_dimensions['A'].width = 28
info.column_dimensions['B'].width = 100
for row in info.iter_rows():
    for cell in row:
        cell.alignment = Alignment(vertical='top', wrap_text=True)

wb.save(out)
print(f'Generated {out} with {len(rows)} dataset cards.')
